"""Build the human title scan of round-2 title/abstract exclusions.

Round 2's title/abstract stage excluded records on the two model passes alone.
The earlier review checked the same stage by scanning the titles of every
model-excluded record for false negatives; this builds the sheet for the same
check, so both rounds are screened by one procedure.

The reviewer reads titles and marks only records that look potentially
eligible (`recheck`). Marked records then get an abstract-level decision; any
that pass go to full text like every other candidate. Nothing is pre-marked:
a suggested answer would invite rubber-stamping.

Ordering is for speed, not for judgement: titles naming a legged platform come
first, because a false negative is most likely among them; the rest follow
alphabetically. The abstract is in the sheet for the marked rows.

Output: work/ta_exclusion_scan.xlsx (refuses to overwrite).
"""
from __future__ import annotations

import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from _common import WORK_DIR

LEGGED = re.compile(r"(?i)quadruped|legged|four-legged|robot dog|robotic dog|spot\b|anymal|unitree|"
                    r"aliengo|go1|go2|b1\b|b2\b|jueying|deep robotics|ghost robotics|vision 60")


def main() -> None:
    out = WORK_DIR / "ta_exclusion_scan.xlsx"
    if out.exists():
        raise SystemExit(f"ERROR: {out} exists; it may already hold decisions.")
    s1 = pd.read_excel(WORK_DIR / "r2_04_screening_round1.xlsx")
    s2 = pd.read_excel(WORK_DIR / "r2_05_screening_round2.xlsx")
    a = s1[s1["Decision round 1 (gpt-oss:120b)"] == "N"].assign(excluded_at="pass 1", model_reason="")
    b = s2[s2["Decision (gpt-oss:120b)"] == "Exclude"].assign(excluded_at="pass 2")
    b["model_reason"] = b["Reason"].fillna("")
    cols = ["record_id", "Title", "Year", "Journal", "excluded_at", "model_reason", "Abstract"]
    d = pd.concat([a[cols], b[cols]], ignore_index=True)
    d["names_legged_platform"] = d["Title"].astype(str).str.contains(LEGGED)
    d = d.sort_values(["names_legged_platform", "Title"], ascending=[False, True])
    d.insert(1, "recheck", "")
    d.insert(2, "note", "")
    d.to_excel(out, index=False)

    wb = load_workbook(out)
    ws = wb.active
    widths = {"A": 10, "B": 9, "C": 24, "D": 80, "E": 6, "F": 30, "G": 9, "H": 30, "I": 60, "J": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"recheck"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{ws.max_row}")
    fill = PatternFill("solid", fgColor="FFF4CC")
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).fill = fill
        ws.cell(r, 4).alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)
    n_leg = int(d["names_legged_platform"].sum())
    print(f"written {out}: {len(d)} excluded records, {n_leg} name a legged platform in the title (listed first)")


if __name__ == "__main__":
    main()
