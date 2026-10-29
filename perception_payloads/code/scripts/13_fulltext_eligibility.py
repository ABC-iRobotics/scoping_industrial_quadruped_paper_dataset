"""
Stage 13 - full-text eligibility screening for round 2, LLM-assisted.

The protocol (osf_protocol.md section 8) makes screening two-stage: title and
abstract, then full text, LLM-assisted with human adjudication, and the human
makes every final inclusion decision. Round-2 records have passed only the
first stage. This script produces the LLM recommendation for the second stage
and the workbook in which P.Z. records the final decision for every record.
The model output is advice; nothing here includes or excludes a study.

DESIGN: EVIDENCE PER CHUNK, THEN ONE DECISION CALL
    Eligibility is a holistic judgement over the whole document, so a verdict
    per chunk cannot simply be merged the way RCI positives are ("any chunk
    wins"): a related-work chunk mentioning a wheeled robot says nothing about
    this study's platform. So the stage runs in two steps:

    1. Every chunk of the full text (30 000 chars, 4 000 overlap, as in
       stage 10) is read with prompts/fulltext_eligibility_evidence.md, which
       asks ten evidence questions mapped onto the PCC criteria and returns a
       verbatim quote for every positive. Quotes are verified mechanically
       against their chunk with stage 10's verify_quote().
    2. One decision call (prompts/fulltext_eligibility.md) receives the
       bibliographic record, the title and abstract, and the collected
       evidence - up to MAX_QUOTES_PER_QUESTION quotes per question,
       verbatim ones first - and applies the PCC criteria from
       osf_protocol.md section 6 verbatim. It returns include / exclude /
       unclear, one reason from a fixed list, a supporting quote and a
       one-line rationale. The decision quote is verified against the full
       text plus the record.

    This reads the whole document, which a fixed prefix would not, and keeps
    the decision call small enough to see all the evidence at once.

GUARDS - no model call is made for these, and the human decides directly
    * unreadable: the amendment-A2 characters-per-page check (a scanned PDF
      would otherwise return a confident exclusion built on nothing).
    * too_long: more than MAX_PAGES pages. A primary research paper in this
      corpus is under 60 pages; a longer file is almost certainly proceedings
      front matter or a whole volume (r2-0630 is a 310-page conference digest).
    * incomplete: at least one chunk or the decision call failed. Re-run to
      retry; never read its decision as final.

KNOWN COMPANION CANDIDATES
    Title/author near-misses against the baseline, found on 26 September 2026
    by fuzzy title matching. They are passed to the decision call as a note and
    flagged in the workbook so the human checks them against the baseline paper.

USAGE
    python scripts/13_fulltext_eligibility.py --pass A
    LLM_MODEL=qwen3.6:35b-a3b python scripts/13_fulltext_eligibility.py --pass B
    python scripts/13_fulltext_eligibility.py --build-sheet

OUTPUTS
    work/fte_<pass>_records.xlsx    one row per record: decision, reason, quote
    work/fte_<pass>_raw/<id>.json   raw chunk and decision replies, for audit
    work/fulltext_decisions.xlsx    human decision workbook (--build-sheet)
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import time
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

from _common import (
    API_URL,
    MODEL,
    WORK_DIR,
    call_llm,
    cell_text,
    chunk_text,
    load_prompt,
    parse_json_reply,
    pdf_page_count,
    pdf_text,
    text_layer_ok,
)

# Stage 10's quote verifier, reused rather than copied. The module name starts
# with a digit, so it is loaded by path.
_spec = importlib.util.spec_from_file_location(
    "rci_extraction", Path(__file__).with_name("10_rci_extraction.py"))
_rci = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_rci)
verify_quote = _rci.verify_quote

EVIDENCE_QUESTIONS = [
    "quadruped_platform",
    "other_platform",
    "industrial_task",
    "industrial_setting",
    "non_industrial_domain",
    "hardware_validation",
    "simulation_only",
    "component_level_only",
    "hardware_configuration",
    "excluded_document_type",
]

DECISIONS = {"include", "exclude", "unclear"}
REASONS = {
    "non_quadruped_platform",
    "no_industrial_use_case",
    "low_level_control_only",
    "simulation_only",
    "not_primary_study_or_excluded_type",
    "non_english",
    "duplicate_or_companion_report",
    "other",
}

COMPANION_CANDIDATES = {
    "r2-0771": "Possible companion of baseline paper009 (same first author J. Li; "
               "'Localized Coverage Planning for a Heat Transfer Tube Inspection Robot', "
               "IEEE RA-L 2025). Check whether it is the same work.",
    "r2-1419": "Possible companion of baseline paper080 ('Evaluating the Robustness of "
               "Autonomous Inspections in the Energy Industry with a Quadruped Robot', "
               "SPE 2023). Check whether it is the same work.",
}

DEFAULT_CHUNK_CHARS = 30_000
DEFAULT_OVERLAP = 4_000
MAX_PAGES = 60
MAX_QUOTES_PER_QUESTION = 3
INPUT_DEFAULT = WORK_DIR / "r2_06_fulltext_needed.xlsx"
PDF_DIR_DEFAULT = WORK_DIR / "pdf_round2"


def collect_evidence(results: list[tuple[int, dict]]) -> tuple[dict, list[str]]:
    """Group verified quotes per question, verbatim first, then by position."""
    rank = {"exact": 0, "elided": 1, "unverified": 2}
    evidence: dict[str, list[dict]] = {}
    for question in EVIDENCE_QUESTIONS:
        found = []
        for offset, payload in results:
            entry = payload.get(question)
            if isinstance(entry, dict) and entry.get("found") is True:
                quote = str(entry.get("quote") or "").strip()
                if quote:
                    found.append({"offset": offset, "quote": quote,
                                  "verified": entry.get("__verified", "unverified")})
        found.sort(key=lambda e: (rank.get(e["verified"], 3), e["offset"]))
        evidence[question] = found[:MAX_QUOTES_PER_QUESTION]
    languages = sorted({str(p.get("language") or "").strip() for _, p in results} - {""})
    return evidence, languages


def decision_input(row: pd.Series, evidence: dict, languages: list[str], note: str) -> str:
    lines = [
        "## Bibliographic record",
        f"Title: {cell_text(row.get('Title'))}",
        f"Year: {cell_text(row.get('Year'))}",
        f"Venue: {cell_text(row.get('Journal'))}",
        f"Publication type (database): {cell_text(row.get('Type'))}",
        f"Abstract: {cell_text(row.get('Abstract'))}",
        "",
        f"## Language(s) detected in the full text: {', '.join(languages) or 'not reported'}",
        "",
        "## Evidence collected from the full text",
    ]
    for question in EVIDENCE_QUESTIONS:
        quotes = evidence.get(question) or []
        if not quotes:
            lines.append(f"- {question}: not found anywhere in the full text")
        else:
            lines.append(f"- {question}:")
            for e in quotes:
                lines.append(f'    - "{e["quote"]}"')
    lines.append("")
    lines.append(f"## Note on companion reports: {note or 'none - do not use duplicate_or_companion_report.'}")
    return "\n".join(lines)


def process_record(row: pd.Series, pdf_path: Path, prompts: dict,
                   chunk_chars: int, overlap: int, delay: float) -> tuple[dict, dict]:
    rid = cell_text(row.get("record_id"))
    pages = pdf_page_count(pdf_path)
    base = {"doc_pages": pages}
    if pages > MAX_PAGES:
        print(f"      {pages} pages > {MAX_PAGES}: not a single paper? Not screened.")
        return {**base, "status": "too_long", "decision": "", "exclusion_reason": ""}, {}

    text = pdf_text(pdf_path)
    usable, ratio = text_layer_ok(text, pages)
    base.update({"doc_chars": len(text), "chars_per_page": round(ratio, 1)})
    if not usable:
        print(f"      TEXT LAYER TOO THIN: {ratio:.0f} chars/page. Not screened.")
        return {**base, "status": "unreadable", "decision": "", "exclusion_reason": ""}, {}

    ev_system, ev_template, ev_params = prompts["evidence"]
    chunks = chunk_text(text, chunk_chars, overlap)
    results, raw_chunks, failed = [], [], 0
    for index, (offset, chunk) in enumerate(chunks, start=1):
        answer = call_llm(ev_system, ev_template.replace("{record}", chunk), ev_params, timeout=300)
        parsed = parse_json_reply(answer)
        raw_chunks.append({"chunk_index": index, "chunk_offset": offset,
                           "chunk_chars": len(chunk), "parsed": parsed is not None, "reply": answer})
        if parsed is None:
            failed += 1
            print(f"      chunk {index}/{len(chunks)}: UNPARSEABLE")
        else:
            for q in EVIDENCE_QUESTIONS:
                entry = parsed.get(q)
                if isinstance(entry, dict) and entry.get("found") is True:
                    entry["__verified"] = verify_quote(entry.get("quote", ""), chunk)
            results.append((offset, parsed))
            hits = [q for q in EVIDENCE_QUESTIONS
                    if isinstance(parsed.get(q), dict) and parsed[q].get("found") is True]
            print(f"      chunk {index}/{len(chunks)}: {len(hits)} signal(s)")
        time.sleep(delay)

    evidence, languages = collect_evidence(results)
    note = COMPANION_CANDIDATES.get(rid, "")
    dec_system, dec_template, dec_params = prompts["decision"]
    user = dec_template.replace("{record}", decision_input(row, evidence, languages, note))
    answer = call_llm(dec_system, user, dec_params, timeout=300)
    parsed = parse_json_reply(answer) or {}

    decision = str(parsed.get("decision") or "").strip().lower()
    reason = str(parsed.get("exclusion_reason") or "").strip()
    quote = str(parsed.get("quote") or "").strip()
    valid = decision in DECISIONS and (decision != "exclude" or reason in REASONS)
    source = text + "\n" + cell_text(row.get("Title")) + "\n" + cell_text(row.get("Abstract"))

    out = {
        **base,
        "status": "ok" if failed == 0 and results and valid else "incomplete",
        "n_chunks": len(chunks),
        "n_failed": failed + (0 if valid else 1),
        "decision": decision if valid else "",
        "exclusion_reason": reason if decision == "exclude" else "",
        "quote": quote,
        "quote_verified": verify_quote(quote, source) if quote else "",
        "rationale": str(parsed.get("rationale") or "").strip(),
        "languages": ", ".join(languages),
        "companion_note": note,
    }
    for q in EVIDENCE_QUESTIONS:
        out[f"ev_{q}"] = len(evidence.get(q) or [])
    raw = {"chunks": raw_chunks, "evidence": evidence, "decision_input": user,
           "decision_reply": answer}
    return out, raw


def run_pass(args) -> None:
    tag = args.pass_label
    out_records = WORK_DIR / f"fte_{tag}_records.xlsx"
    raw_dir = WORK_DIR / f"fte_{tag}_raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    pdf_dir = Path(args.pdf_dir)

    records = pd.read_excel(args.input)
    if out_records.exists():
        frame = pd.read_excel(out_records, dtype=object)
    else:
        frame = records[["record_id", "Title", "Year", "Doi"]].copy().astype(object)
    if "status" not in frame.columns:
        frame["status"] = pd.Series([pd.NA] * len(frame), dtype=object)
    frame = frame.set_index("record_id")
    rec = records.set_index("record_id")

    prompts = {"evidence": load_prompt("fulltext_eligibility_evidence"),
               "decision": load_prompt("fulltext_eligibility")}
    print(f"pass={tag}  model={MODEL}  endpoint={API_URL}")

    processed = 0
    for position, rid in enumerate(frame.index, start=1):
        if cell_text(frame.at[rid, "status"]) in {"ok", "too_long", "unreadable"}:
            continue
        pdf_path = pdf_dir / f"{rid}.pdf"
        if not pdf_path.exists():
            frame.at[rid, "status"] = "pdf_missing"
            continue
        if args.limit and processed >= args.limit:
            break
        print(f"  {position}/{len(frame)}  {rid}")
        out, raw = process_record(rec.loc[rid], pdf_path, prompts,
                                  args.chunk_chars, args.overlap, args.delay)
        for key, value in out.items():
            if key not in frame.columns:
                frame[key] = pd.Series([pd.NA] * len(frame), index=frame.index, dtype=object)
            frame.at[rid, key] = value
        (raw_dir / f"{rid}.json").write_text(json.dumps({
            "record_id": rid, "pass": tag, "model": MODEL, "endpoint": API_URL,
            "chunk_chars": args.chunk_chars, "overlap": args.overlap,
            "screened_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            **raw,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"    -> {out['status']} {out.get('decision', '')} {out.get('exclusion_reason', '')}\n")
        processed += 1
        frame.reset_index().to_excel(out_records, index=False)

    frame.reset_index().to_excel(out_records, index=False)
    print(f"\nProcessed {processed} record(s). Status: {frame['status'].astype(str).value_counts().to_dict()}")
    ok = frame[frame["status"].astype(str) == "ok"]
    if len(ok):
        print("Decisions:", ok["decision"].value_counts().to_dict())
        print("Reasons:", ok["exclusion_reason"].replace("", pd.NA).dropna().value_counts().to_dict())
    print(f"Records: {out_records}")


def build_sheet(args) -> None:
    out = WORK_DIR / "fulltext_decisions.xlsx"
    if out.exists():
        raise SystemExit(f"ERROR: {out} exists and may hold human decisions. Refusing to overwrite.")
    records = pd.read_excel(args.input)[["record_id", "Title", "Year", "Doi"]]
    sheet = records.rename(columns={"Title": "title", "Year": "year", "Doi": "doi"})
    cols = ["status", "decision", "exclusion_reason", "quote", "quote_verified", "rationale"]
    for tag in ("A", "B"):
        path = WORK_DIR / f"fte_{tag}_records.xlsx"
        if not path.exists():
            continue
        p = pd.read_excel(path, dtype=object)[["record_id"] + [c for c in cols]]
        sheet = sheet.merge(p.rename(columns={c: f"{tag}_{c}" for c in cols}), on="record_id", how="left")
    if "A_decision" in sheet and "B_decision" in sheet:
        a = sheet["A_decision"].fillna("").astype(str)
        b = sheet["B_decision"].fillna("").astype(str)
        ar = sheet["A_exclusion_reason"].fillna("").astype(str)
        br = sheet["B_exclusion_reason"].fillna("").astype(str)
        sheet["passes_agree"] = [
            "" if not x or not y else ("yes" if x == y and xr == yr else ("decision only" if x == y else "no"))
            for x, y, xr, yr in zip(a, b, ar, br)]
    sheet["companion_check"] = sheet["record_id"].map(COMPANION_CANDIDATES).fillna("")
    sheet["final_decision"] = ""
    sheet["final_reason"] = ""
    sheet["note"] = ""
    sheet.to_excel(out, index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = load_workbook(out)
    ws = wb.active
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "C2"
    header = [c.value for c in ws[1]]
    fill = PatternFill("solid", fgColor="FFF4CC")
    n = ws.max_row
    for name, allowed in (("final_decision", "include,exclude"),
                          ("final_reason", ",".join(sorted(REASONS)))):
        col = ws.cell(row=1, column=header.index(name) + 1).column_letter
        dv = DataValidation(type="list", formula1=f'"{allowed}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
        for r in range(2, n + 1):
            ws[f"{col}{r}"].fill = fill
    wb.save(out)
    print(f"Wrote {out} ({len(sheet)} rows).")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    parser.add_argument("--input", default=str(INPUT_DEFAULT))
    parser.add_argument("--pdf-dir", default=str(PDF_DIR_DEFAULT))
    parser.add_argument("--pass", dest="pass_label", default="A")
    parser.add_argument("--chunk-chars", type=int, default=DEFAULT_CHUNK_CHARS)
    parser.add_argument("--overlap", type=int, default=DEFAULT_OVERLAP)
    parser.add_argument("--delay", type=float, default=2.0)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--build-sheet", action="store_true",
                        help="build the human decision workbook from the pass outputs")
    args = parser.parse_args()
    if args.build_sheet:
        build_sheet(args)
    else:
        run_pass(args)


if __name__ == "__main__":
    main()
