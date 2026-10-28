"""Rebuild a blind-validation coding workbook in the v2 layout.

The v1 layout (written by 32_validation_sample.py) followed the order of the
two extraction prompts, carried a quote column after every field, and grew two
columns at the far end when the holistic reproducibility judgement and the
AI/ML field were added. P.Z. asked for a cleaner sheet on 5 Oct 2026.

v2 groups the columns in the order a coder meets the information in a paper,
drops the quote columns that free-text and categorical fields never needed,
and drops `minutes_spent`, which is measured from ActivityWatch instead.

Nothing is lost: every value already entered is carried across, and the whole
v1 sheet is kept as a hidden `v1_archive` sheet in the new file, so the quotes
entered in the dropped columns survive.

Usage:
    python scripts/33_validation_sheet_v2.py work/validation/validation_coding_baseline.xlsx
writes work/validation/validation_coding_baseline_v2.xlsx; refuses to overwrite.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

YN = ["yes", "no"]
SPEC = ["model", "modality", "unspecified"]
KIND = ["real", "simulated", "both", "unclear"]

# (group, column, allowed values or None, width, rule)
LAYOUT = [
    ("Study", "record_id", None, 11, ""),
    ("Study", "title", None, 40, ""),
    ("Study", "year", None, 6, ""),
    ("Study", "doi", None, 18, ""),

    ("Platform & payload", "rci1_platform_model", YN, 10, "Robot platform named at model level (e.g. 'Boston Dynamics Spot', 'ANYmal C'). 'a quadruped robot' = no."),
    ("Platform & payload", "rci1_platform_model__quote", None, 30, "Verbatim sentence for a yes."),
    ("Platform & payload", "optical_sensors_raw", None, 26, "Every optical sensor as written in the paper, comma separated - class-only mentions too (e.g. 'RGB-D camera', 'a LiDAR'), not only named models. Empty only if no optical sensor is mentioned."),
    ("Platform & payload", "sensor_specificity", SPEC, 12, "Most specific naming of any optical sensor: model / modality (class only) / unspecified."),
    ("Platform & payload", "rci2_sensors_model", YN, 10, "At least one optical sensor named at make-and-model level (e.g. 'Intel RealSense D455'). 'an RGB-D camera' = no."),
    ("Platform & payload", "rci2_sensors_model__quote", None, 30, "Verbatim sentence for a yes."),
    ("Platform & payload", "onboard_computer_raw", None, 22, "Onboard computing hardware as written. Empty if not stated."),
    ("Platform & payload", "compute_specificity", SPEC, 12, "Same three levels for the onboard computer."),
    ("Platform & payload", "rci3_compute_model", YN, 10, "Onboard computer named at model level AND it is the computer the ROBOT CARRIES. A workstation/laptop/server running experiments = no. When in doubt = no."),
    ("Platform & payload", "rci3_compute_model__quote", None, 30, "Verbatim sentence for a yes."),

    ("Software & methods", "rci4_middleware_version", YN, 10, "Middleware/framework named WITH version (e.g. 'ROS 2 Humble', 'Ubuntu 20.04', 'PyTorch 1.13'). 'ROS' alone = no."),
    ("Software & methods", "rci4_middleware_version__quote", None, 30, "Verbatim sentence for a yes."),
    ("Software & methods", "rci5_algorithms_cited", YN, 10, "Core perception/navigation algorithm named AND attributed: citation or specific named method (e.g. 'LIO-SAM [12]', 'YOLOv8'). 'a deep-learning detector' = no."),
    ("Software & methods", "rci5_algorithms_cited__quote", None, 30, "Verbatim sentence for a yes."),
    ("Software & methods", "ai_ml_solutions_raw", None, 22, "ML/AI method NAMES used for perception (e.g. 'YOLOv8, PointNet++'). Names only, not a sentence. Empty if none."),
    ("Software & methods", "rci6_parameters", YN, 10, "At least one concrete operating parameter or configuration value (e.g. '10 Hz', 'voxel size 0.4 m', '640x480', 'FP16')."),
    ("Software & methods", "rci6_parameters__quote", None, 30, "Verbatim sentence for a yes. One is enough."),

    ("Evaluation", "eval_environment_kind", KIND, 12, "Where the evaluation took place: real / simulated / both / unclear."),
    ("Evaluation", "rci7_eval_environment", YN, 10, "The TEST SETTING has at least one measurable quantity: dimension, distance, duration of the trial, or count of assets/obstacles. A measured RESULT is rci8, not this. 'a real substation' = no. Real vs simulated does not matter here."),
    ("Evaluation", "rci7_eval_environment__quote", None, 30, "Verbatim sentence for a yes."),
    ("Evaluation", "cap_real_world_validation", YN, 10, "Integrated system demonstrated in a physical real-world setting, not only simulation or lab floor."),
    ("Evaluation", "cap_real_world_validation__quote", None, 30, "Verbatim sentence showing the physical trial."),
    ("Evaluation", "rci8_success_metric", YN, 10, "A QUANTITATIVE result for the goal task (accuracy, success rate, RMSE, time...). 'performed well' = no."),
    ("Evaluation", "rci8_success_metric__quote", None, 30, "Verbatim sentence for a yes."),

    ("Capabilities", "cap_slam_mapping", YN, 10, "System builds or updates a map (SLAM, mapping). Only USING a pre-built map = no."),
    ("Capabilities", "cap_slam_mapping__quote", None, 26, "Verbatim sentence for a yes."),
    ("Capabilities", "cap_autonomous_navigation", YN, 10, "Robot plans and follows paths without a human driving. Shared/supervised autonomy = yes (note it). Teleoperation only = no."),
    ("Capabilities", "cap_autonomous_navigation__quote", None, 26, "Verbatim sentence for a yes."),
    ("Capabilities", "cap_obstacle_avoidance", YN, 10, "System detects and avoids / traverses obstacles as a described function."),
    ("Capabilities", "cap_obstacle_avoidance__quote", None, 26, "Verbatim sentence for a yes."),
    ("Capabilities", "cap_object_detection", YN, 10, "System detects, recognises, segments or classifies objects, defects or people."),
    ("Capabilities", "cap_object_detection__quote", None, 26, "Verbatim sentence for a yes."),
    ("Capabilities", "cap_instrument_reading", YN, 10, "System reads gauges, dials, displays, meters or signage."),
    ("Capabilities", "cap_instrument_reading__quote", None, 26, "Verbatim sentence for a yes."),
    ("Capabilities", "cap_target_tracking", YN, 10, "System tracks a specific moving target or person over time."),
    ("Capabilities", "cap_target_tracking__quote", None, 26, "Verbatim sentence for a yes."),

    ("Context", "industry_sector_raw", None, 22, "Industrial sector/setting in the paper's own words (short)."),
    ("Context", "high_level_goal_raw", None, 30, "Primary objective in the authors' own words (short)."),

    ("Openness & overall", "rci9_artifact", YN, 10, "Code, data or model PRODUCED BY THESE AUTHORS is publicly available. Someone else's tool they used = no. Publisher 'supplementary material available at <doi>' boilerplate = no, unless it says the supplement is code/data/model."),
    ("Openness & overall", "rci9_artifact__quote", None, 26, "Verbatim sentence for a yes."),
    ("Openness & overall", "reproducible_holistic", YN, 10, "Round-1 holistic question: 'Is the work reproducible from the DESCRIPTION (sufficient details on hardware/software)?' Judge the text, not code availability - code is rci9. Overall judgement after reading."),
    ("Openness & overall", "reproducible_holistic__note", None, 30, "One line: what is or is not described well enough."),

    ("Notes", "notes", None, 40, "Anything unclear, borderline calls, decisions you want to remember."),
]

GROUP_FILL = {
    "Study": "D9D9D9",
    "Platform & payload": "C6E0F5",
    "Software & methods": "D5E8D4",
    "Evaluation": "FCE4D6",
    "Capabilities": "E4DFEC",
    "Context": "FFF2CC",
    "Openness & overall": "F8CBAD",
    "Notes": "D9D9D9",
}
INPUT_FILL = PatternFill("solid", fgColor="FFFBEA")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    dst = src.with_name(src.stem + "_v2.xlsx")
    if dst.exists():
        raise SystemExit(f"ERROR: {dst} exists; refusing to overwrite a coding sheet.")

    old = load_workbook(src)
    ws_old = old["coding"]
    head_old = [c.value for c in ws_old[1]]
    # v1 headers carry suffixes such as " (NOT NEEDED)"; match on the bare name.
    index = {str(h).split(" (")[0]: i for i, h in enumerate(head_old) if h}
    rows = [[c.value for c in r] for r in ws_old.iter_rows(min_row=2)]

    used_old = {name for _, name, _, _, _ in LAYOUT}
    dropped_with_data = sorted(
        name for name, i in index.items()
        if name not in used_old and any(r[i] not in (None, "") for r in rows)
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "coding"
    for j, (group, name, allowed, width, _) in enumerate(LAYOUT, start=1):
        cell = ws.cell(1, j, name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor=GROUP_FILL[group])
        ws.column_dimensions[cell.column_letter].width = width
        for i, row in enumerate(rows, start=2):
            value = row[index[name]] if name in index else None
            c = ws.cell(i, j, value)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if group != "Study":
                c.fill = INPUT_FILL
        if allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{cell.column_letter}2:{cell.column_letter}{len(rows) + 1}")
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"

    rules = wb.create_sheet("rules")
    rules.append(["group", "column", "allowed values", "rule"])
    for c in rules[1]:
        c.font = Font(bold=True)
    for group, name, allowed, _, rule in LAYOUT:
        if rule:
            rules.append([group, name, ", ".join(allowed) if allowed else "free text", rule])
    rules.append([])
    rules.append(["", "", "", "Code from the full text only. Do not open any adjudication file, codebook, or extraction output until every row here is finished."])
    rules.append(["", "", "", "Quotes only for yes answers of yes/no items. 'no' needs no quote. Free-text fields: short text in the paper's words."])
    rules.append(["", "", "", "Time per study is measured from ActivityWatch; no minutes column."])
    for col, w in zip("ABCD", (20, 32, 22, 120)):
        rules.column_dimensions[col].width = w

    archive = wb.create_sheet("v1_archive")
    archive.append(head_old)
    for r in rows:
        archive.append(r)
    archive.sheet_state = "hidden"

    wb.save(dst)
    print(f"written {dst}: {len(rows)} rows, {len(LAYOUT)} columns")
    print("v1 columns dropped that held data (kept in hidden v1_archive):", dropped_with_data or "none")


if __name__ == "__main__":
    main()
