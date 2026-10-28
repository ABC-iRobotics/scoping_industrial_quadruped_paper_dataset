"""Draw the blind human-validation subsample and build the coding workbook.

The subsample is the human reference standard against which LLM extraction
accuracy is measured (CLAUDE.md section 6, component 4). Two properties make it
valid, and this script exists to enforce both:

1. The draw is random, seeded and recorded before anyone looks at model output
   for the drawn studies. The seed and the drawn ids are written to
   ``protocol/validation_sample.md``, which is committed, so the draw can be
   re-run and checked.
2. The coder never sees model output. The workbook carries only identifiers,
   titles and the coding rules - no pass-A/B codes, no round-1 charting values.
   The input list deliberately keeps only ``Pdf name``, ``Title``, ``Year`` and
   ``Doi``, because ``rci_include_list.xlsx`` also carries the round-1 LLM
   ``Optical Sensors`` / ``Onboard Computer`` strings.

The draw has two strata. The baseline stratum (round 1) is drawn now. The
round-2 stratum is drawn once the round-2 full texts are in, with the same seed
and ``--stratum round2``; round-2 extraction has not run yet, so nothing in it
can be contaminated.

Usage:
    python scripts/32_validation_sample.py --stratum baseline --n 24
    python scripts/32_validation_sample.py --stratum round2 --n 6 --ids-from <xlsx>
"""
from __future__ import annotations

import argparse
import random
import shutil
from datetime import date
from pathlib import Path

import fitz  # PyMuPDF
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from _common import PDF_DIR

HERE = Path(__file__).resolve().parent.parent
WORK = HERE / "work"
OUT = WORK / "validation"
RECORD = HERE / "protocol" / "validation_sample.md"

SEED = 20260926

# Studies whose model output was read during the five-study calibration.
# A coder who has seen their outputs cannot code them blind.
CONTAMINATED = {"paper031", "paper033", "paper041", "paper043", "paper056"}

# Same guard as amendment A2: below this, the PDF has no usable text layer.
MIN_CHARS_PER_PAGE = 1000

YN = ["yes", "no"]
SPEC = ["model", "modality", "unspecified"]
KIND = ["real", "simulated", "both", "unclear"]

# (column, allowed values or None for free text, rule shown to the coder)
ITEMS = [
    ("rci1_platform_model", YN, "Robot platform named at model level (e.g. 'Boston Dynamics Spot', 'ANYmal C'). 'a quadruped robot' = no."),
    ("rci2_sensors_model", YN, "At least one optical sensor named at make-and-model level (e.g. 'Intel RealSense D455'). 'an RGB-D camera' = no."),
    ("rci3_compute_model", YN, "Onboard computer named at model level AND it is the computer the ROBOT CARRIES. A workstation/laptop/server running experiments = no. When in doubt = no."),
    ("rci4_middleware_version", YN, "Middleware/framework named WITH version (e.g. 'ROS 2 Humble', 'Ubuntu 20.04', 'PyTorch 1.13'). 'ROS' alone = no."),
    ("rci5_algorithms_cited", YN, "Core perception/navigation algorithm named AND attributed: citation or specific named method (e.g. 'LIO-SAM [12]', 'YOLOv8'). 'a deep-learning detector' = no."),
    ("rci6_parameters", YN, "At least one concrete operating parameter or configuration value (e.g. '10 Hz', 'voxel size 0.4 m', '640x480')."),
    ("rci7_eval_environment", YN, "Evaluation environment has at least one MEASURABLE QUANTITY: dimension, distance, duration or count of assets. 'a real substation' = no. Real vs simulated does not matter here."),
    ("rci8_success_metric", YN, "A QUANTITATIVE result for the goal task (accuracy, success rate, RMSE, time...). 'performed well' = no."),
    ("rci9_artifact", YN, "Code, data or model PRODUCED BY THESE AUTHORS is publicly available. Someone else's tool they used = no. Publisher 'supplementary material available at <doi>' boilerplate = no, unless it says the supplement is code/data/model."),
    ("optical_sensors_raw", None, "Every optical sensor as written in the paper, comma separated. Empty if none."),
    ("onboard_computer_raw", None, "Onboard computing hardware as written. Empty if not stated."),
    ("sensor_specificity", SPEC, "Most specific naming of any optical sensor: model / modality (class only) / unspecified."),
    ("compute_specificity", SPEC, "Same three levels for the onboard computer."),
    ("eval_environment_kind", KIND, "Where the evaluation took place: real / simulated / both / unclear."),
    ("cap_slam_mapping", YN, "System builds or updates a map (SLAM, mapping). Only USING a pre-built map = no."),
    ("cap_autonomous_navigation", YN, "Robot plans and follows paths without a human driving. Shared/supervised autonomy = yes (note it). Teleoperation only = no."),
    ("cap_obstacle_avoidance", YN, "System detects and avoids / traverses obstacles as a described function."),
    ("cap_object_detection", YN, "System detects, recognises, segments or classifies objects, defects or people."),
    ("cap_instrument_reading", YN, "System reads gauges, dials, displays, meters or signage."),
    ("cap_target_tracking", YN, "System tracks a specific moving target or person over time."),
    ("cap_real_world_validation", YN, "Integrated system demonstrated in a physical real-world setting, not only simulation or lab floor."),
    ("industry_sector_raw", None, "Industrial sector/setting in the paper's own words."),
    ("high_level_goal_raw", None, "Primary objective in the authors' own words."),
    # Added 29 Sep 2026 after the first baseline study was coded. The protocol
    # (section 12) requires the round-1 holistic judgement to be validated in the
    # subsample, and the capability instrument extracts AI/ML methods.
    ("ai_ml_solutions_raw", None, "ML/AI methods used for perception, as named in the paper (e.g. 'YOLOv8', 'PointNet++'). Empty if none."),
    ("reproducible_holistic", YN, "Round-1 holistic question: 'Is the work reproducible from the description (sufficient details on hardware/software)?' Overall judgement, independent of the RCI items."),
]

# Only yes/no items need an evidence quote; free-text and categorical fields
# carry their own evidence.
NEEDS_QUOTE = {col for col, allowed, _ in ITEMS if allowed == YN}


def text_ok(pdf: Path) -> tuple[bool, int]:
    doc = fitz.open(pdf)
    chars = sum(len(p.get_text()) for p in doc)
    per_page = chars // max(doc.page_count, 1)
    return per_page >= MIN_CHARS_PER_PAGE, per_page


def baseline_frame(pdf_dir: Path) -> pd.DataFrame:
    df = pd.read_excel(WORK / "rci_include_list.xlsx", usecols=["Pdf name", "Title", "Year", "Doi"])
    df = df.rename(columns={"Pdf name": "record_id"})
    df["pdf"] = df.record_id.map(lambda r: pdf_dir / f"{r}.pdf")
    return df


def build_workbook(sample: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "coding"
    head = ["record_id", "title", "year", "doi"]
    for col, _, _ in ITEMS:
        head += [col, f"{col}__quote" if col in NEEDS_QUOTE else f"{col}__note"]
    head += ["notes"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for _, r in sample.iterrows():
        ws.append([r.record_id, r.Title, int(r.Year), r.Doi])
    ws.freeze_panes = "E2"
    fill = PatternFill("solid", fgColor="FFF4CC")
    for i, (col, allowed, _) in enumerate(ITEMS):
        cidx = 5 + 2 * i
        letter = ws.cell(row=1, column=cidx).column_letter
        ws.column_dimensions[letter].width = 14
        ws.column_dimensions[ws.cell(row=1, column=cidx + 1).column_letter].width = 30
        if allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{len(sample) + 1}")
        for row in range(2, len(sample) + 2):
            ws.cell(row=row, column=cidx).fill = fill
    ws.column_dimensions["B"].width = 50

    rules = wb.create_sheet("rules")
    rules.append(["column", "allowed values", "rule"])
    for c in rules[1]:
        c.font = Font(bold=True)
    for col, allowed, rule in ITEMS:
        rules.append([col, ", ".join(allowed) if allowed else "free text", rule])
    rules.column_dimensions["A"].width = 28
    rules.column_dimensions["B"].width = 26
    rules.column_dimensions["C"].width = 120
    rules.append([])
    rules.append(["", "", "Code from the full text only. Do not open any adjudication file, codebook, or extraction output until every row here is finished."])
    rules.append(["", "", "For every 'yes' paste the shortest sentence that proves it into the __quote column. 'no' needs no quote."])
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--stratum", choices=["baseline", "round2"], required=True)
    ap.add_argument("--n", type=int, required=True)
    ap.add_argument("--pdf-dir", default=None)
    ap.add_argument("--ids-from", default=None, help="round2: xlsx of included round-2 records (record_id, title, year, doi)")
    args = ap.parse_args()

    OUT.mkdir(parents=True, exist_ok=True)
    wb_path = OUT / f"validation_coding_{args.stratum}.xlsx"
    if wb_path.exists():
        raise SystemExit(f"ERROR: {wb_path} exists. The draw is done once; refusing to overwrite a coding sheet.")

    if args.stratum == "baseline":
        frame = baseline_frame(Path(args.pdf_dir) if args.pdf_dir else PDF_DIR)
    else:
        if not args.ids_from:
            raise SystemExit("ERROR: --ids-from is required for the round2 stratum.")
        frame = pd.read_excel(args.ids_from).rename(columns=str.strip)
        frame = frame.rename(columns={"title": "Title", "year": "Year", "doi": "Doi"})
        pdf_dir = Path(args.pdf_dir) if args.pdf_dir else WORK / "pdf_round2"
        frame["pdf"] = frame.record_id.map(lambda r: pdf_dir / f"{r}.pdf")

    excluded = []
    keep = []
    for _, r in frame.sort_values("record_id").iterrows():
        if r.record_id in CONTAMINATED:
            excluded.append((r.record_id, "calibration: model output already seen"))
        elif not r.pdf.exists():
            excluded.append((r.record_id, "no PDF"))
        else:
            ok, per_page = text_ok(r.pdf)
            if ok:
                keep.append(r)
            else:
                excluded.append((r.record_id, f"no usable text layer ({per_page} chars/page)"))
    eligible = pd.DataFrame(keep)

    rng = random.Random(f"{SEED}-{args.stratum}")
    ids = sorted(rng.sample(sorted(eligible.record_id), args.n))
    sample = eligible[eligible.record_id.isin(ids)].sort_values("record_id")

    pdf_out = OUT / "pdf"
    pdf_out.mkdir(exist_ok=True)
    for _, r in sample.iterrows():
        shutil.copy2(r.pdf, pdf_out / f"{r.record_id}.pdf")
    build_workbook(sample, wb_path)

    with RECORD.open("a", encoding="utf-8") as f:
        if RECORD.stat().st_size == 0:
            f.write("# Blind validation subsample - draw record\n\n"
                    "Written by `scripts/32_validation_sample.py`. Each stratum is drawn once.\n"
                    "The draw is `random.Random(f\"{SEED}-{stratum}\").sample(sorted(eligible), n)`.\n\n")
        f.write(f"## Stratum `{args.stratum}` - drawn {date.today().isoformat()}\n\n")
        f.write(f"- Seed: `{SEED}-{args.stratum}`\n- Frame: {len(frame)} records; eligible {len(eligible)}; drawn {args.n}\n")
        f.write("- Excluded before the draw:\n")
        for rid, why in excluded:
            f.write(f"  - `{rid}` - {why}\n")
        f.write(f"- Drawn: {', '.join(f'`{i}`' for i in ids)}\n\n")

    print(f"frame {len(frame)}, eligible {len(eligible)}, drawn {len(ids)}")
    print("excluded:", excluded)
    print("drawn:", ids)
    print("workbook:", wb_path)


if __name__ == "__main__":
    main()
