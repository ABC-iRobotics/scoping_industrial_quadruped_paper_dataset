"""Merge a partial coding workbook back into the master validation sheet.

P.Z. coded the last baseline studies in a separate workbook holding only those
rows (validation_coding_remaining.xlsx, 8 Oct 2026). This copies every coded
cell of those rows into the master by record_id. It refuses to overwrite a
master cell that already holds a different value, so nothing coded earlier can
be lost silently; the master is backed up to _archive/ before it is written.

Usage:
    python scripts/34_merge_validation.py work/validation/validation_coding_baseline_v2.xlsx \
                                          work/validation/validation_coding_remaining.xlsx
"""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    master_path, part_path = map(Path, sys.argv[1:3])
    master = load_workbook(master_path)
    part = load_workbook(part_path)
    ms, ps = master["coding"], part["coding"]
    mh = {c.value: i + 1 for i, c in enumerate(ms[1])}
    ph = {c.value: i + 1 for i, c in enumerate(ps[1])}
    if set(ph) - set(mh):
        raise SystemExit(f"ERROR: columns missing in master: {set(ph) - set(mh)}")
    mrow = {ms.cell(r, mh["record_id"]).value: r for r in range(2, ms.max_row + 1)}

    copied, conflicts = 0, []
    for r in range(2, ps.max_row + 1):
        rid = ps.cell(r, ph["record_id"]).value
        if rid is None:
            continue
        if rid not in mrow:
            raise SystemExit(f"ERROR: {rid} not in master")
        for name, j in ph.items():
            new = ps.cell(r, j).value
            if new in (None, ""):
                continue
            cell = ms.cell(mrow[rid], mh[name])
            if cell.value not in (None, "") and cell.value != new:
                conflicts.append((rid, name, cell.value, new))
                continue
            if cell.value != new:
                cell.value = new
                copied += 1
    if conflicts:
        for c in conflicts:
            print("CONFLICT", c)
        raise SystemExit("ERROR: conflicting values; master not written.")

    archive = master_path.parent / "_archive"
    archive.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    shutil.copy2(master_path, archive / f"{master_path.stem}.before_merge_{stamp}.xlsx")
    master.save(master_path)
    print(f"merged {copied} cells from {part_path.name} into {master_path.name}")


if __name__ == "__main__":
    main()
